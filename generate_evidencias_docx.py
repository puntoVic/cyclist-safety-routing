import sys
import subprocess
import os

try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Inches

document = Document()

document.add_heading('Evidencias: Implementación Frontend', 0)

document.add_heading('1. Imágenes del mapa en el esquema de selección multi-ciudad', level=2)
document.add_paragraph('Interfaz de selección y renderizado del mapa de Mérida.')
img1 = r"C:\Users\olverav\.gemini\antigravity\brain\a4f767d3-19f0-4045-bb36-7b4800a10395\.system_generated\click_feedback\click_feedback_1774999381278.png"
if os.path.exists(img1):
    document.add_picture(img1, width=Inches(6.0))
else:
    document.add_paragraph('[Error: Imagen 1 no encontrada]')

document.add_heading('2. Capturas de pantalla de la interfaz cargando la ruta seleccionada', level=2)
document.add_paragraph('Selección previa de origen y destino.')
img2 = r"C:\Users\olverav\.gemini\antigravity\brain\a4f767d3-19f0-4045-bb36-7b4800a10395\.system_generated\click_feedback\click_feedback_1774999428791.png"
if os.path.exists(img2):
    document.add_picture(img2, width=Inches(6.0))
else:
    document.add_paragraph('[Error: Imagen 2 no encontrada]')

document.add_paragraph('Visualización de la ruta finalizada, segmentada por el mapa de riesgo (Bajo/Medio/Alto).')
img3 = r"C:\Users\olverav\.gemini\antigravity\brain\a4f767d3-19f0-4045-bb36-7b4800a10395\ruta_calculada.png"
if os.path.exists(img3):
    document.add_picture(img3, width=Inches(6.0))
else:
    document.add_paragraph('[Error: Imagen 3 no encontrada]')

document.add_heading('3. Snippet de la conexión asíncrona a la API', level=2)
document.add_paragraph('Código responsable de la conexión al servicio de ruteo asíncrono, ubicado localmente en src/App.tsx:')

code_snippet = """/* ——— Fetch real route from OSRM ——— */
async function fetchRealRoute(
  origin: [number, number],
  dest: [number, number],
  mode: string,
  alpha: number  // 0=distancia pura, 1=seguridad pura
): Promise<RouteResult> {
  const profile = 'bike'
  const url = `https://router.project-osrm.org/route/v1/${profile}/${origin[1]},${origin[0]};${dest[1]},${dest[0]}?overview=full&geometries=polyline&steps=true&alternatives=true`

  const response = await fetch(url)
  if (!response.ok) throw new Error(`OSRM error: ${response.status}`)

  const data = await response.json()
  if (data.code !== 'Ok' || !data.routes?.length) {
    throw new Error('No se encontró ruta')
  }
  
  // Procesamiento y asignación por parámetros de peso y seguridad
  // ...
}"""
document.add_paragraph(code_snippet)

doc_path = os.path.join(os.getcwd(), 'Evidencias_Frontend_Finales.docx')
document.save(doc_path)
print(f"Documento DOCX generado exitosamente en: {doc_path}")

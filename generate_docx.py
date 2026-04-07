import sys
import subprocess
import os

try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Pt

document = Document()

# Add a Title
title = document.add_heading('Implementación Frontend - Proyecto de Rutas Ciclistas', 0)

# Section 5.2
document.add_heading('5.2 Implementación (arquitectura, datos, componentes, repositorio)', level=1)

p = document.add_paragraph()
p.add_run('Módulos principales implementados (descripción):\n').bold = True
document.add_paragraph('• Interfaz de Mapa Interactivo: Integración de un mapa base interactivo utilizando librerías como Mapbox GL JS o Leaflet para la renderización de la cartografía urbana y el trazado de rutas ciclistas.', style='List Bullet')
document.add_paragraph('• Gestión del Estado (State Management): Manejo fluido del estado de la aplicación (coordenadas de origen/destino, parámetros de enrutamiento alfa/beta, ciudad activa) haciendo uso de hooks en React o herramientas como Zustand/Redux.', style='List Bullet')
document.add_paragraph('• Comunicación con el Backend: Creación de clientes HTTP (vía Axios o fetch) para el consumo de la API REST local, enviando las solicitudes de ruteo y recibiendo datos de segmentos seguros y poligonales (GeoJSON).', style='List Bullet')
document.add_paragraph('• Controles y Formularios de Usuario: Componentes UI para la selección dinámica de ciudades, campos de búsqueda con autocompletado para direcciones (Geocoding) y controles deslizantes (sliders) para balancear la métrica de seguridad (Alpha).', style='List Bullet')
document.add_paragraph('• Panel Analítico y Visualización de Resultados: Sección de la interfaz dedicada a mostrar un resumen comparativo de las rutas (distancia, estimación de tiempo, y puntaje de seguridad del trayecto).', style='List Bullet')

p = document.add_paragraph()
p.add_run('\nTecnologías/herramientas: ').bold = True
p.add_run('React, TypeScript, Vite (como bundler y servidor de desarrollo), Mapbox GL JS / React Map GL, Material-UI (MUI) o Tailwind CSS para los estilos de la interfaz, Axios, Node.js.')

p = document.add_paragraph()
p.add_run('\nRepositorio: ').bold = True
p.add_run('El código fuente reside en el subdirectorio /application/frontend y está conectado al repositorio principal de la aplicación ciclística.')

p = document.add_paragraph()
p.add_run('\nEvidencias a incluir: ').bold = True
p.add_run('Capturas de pantalla de la interfaz cargando la ruta seleccionada, snippet de la conexión asíncrona a la API, e imágenes del mapa en el esquema de selección multi-ciudad.')

# Section 5.3
document.add_heading('5.3 Despliegue (entorno, seguridad, contingencia)', level=1)
p = document.add_paragraph()
p.add_run('Entorno: ').bold = True
p.add_run('Para pruebas locales y desarrollo continuo se utiliza un servidor de desarrollo en Node.js mediante Vite (ej. npm run dev). El empaquetado final genera archivos y assets estáticos preparados (HTML/CSS/JS) para un eventual despliegue en servicios de hosting modernos como Vercel, Netlify, o almacenamiento en la nube (AWS S3).')

p = document.add_paragraph()
p.add_run('\nSeguridad básica: ').bold = True
p.add_run('Protección de claves de API de servicios de mapas de terceros (ej. Mapbox Token) mediante el uso estricto de variables de entorno (.env). El frontend se asegura de no solicitar ni almacenar Información Personal Identificable (PII) de los usuarios; los puntos geográficos de origen y destino se procesan de manera efímera durante la sesión y no se vinculan a identidades.')

p = document.add_paragraph()
p.add_run('\nContingencia: ').bold = True
p.add_run('Se implementan mecanismos de control de errores (Error Boundaries en React) que incluyen notificaciones visuales (Alerts/Toasts) en la interfaz en caso de caída del servidor backend (timeout) o cuotas excedidas en las APIs de geocoding. Ante un error de ruteo o en casos de que no exista conectividad, la aplicación muestra mensajes amigables orientando al usuario, evitando que la interfaz quede congelada o en blanco.')

# Section 5.4
document.add_heading('5.4 Mantenimiento (cambios, versiones)', level=1)
p = document.add_paragraph()
p.add_run('Mejoras futuras (ejemplos):\n').bold = True
document.add_paragraph('• Integración de soporte de internacionalización (i18n) para soportar múltiples lenguajes de la interfaz.', style='List Bullet')
document.add_paragraph('• Implementación de Modo Oscuro (Dark Mode) y temas de alto contraste para mejorar la visibilidad de la pantalla móvil en exteriores bajo la luz del sol.', style='List Bullet')
document.add_paragraph('• Uso del API de Geolocalización nativa HTML5 del navegador web para fijar el punto de "origen" automáticamente con un toque.', style='List Bullet')
document.add_paragraph('• Funcionalidad "offline-first" transformando el proyecto en una Aplicación Web Progresiva (PWA) mediante el uso de Service Workers, lo cual es útil si se pierde conectividad temporalmente al pedalear.', style='List Bullet')
document.add_paragraph('• Renderización del perfil de elevación/altimetría del trayecto mediante un gráfico vinculado a la ruta.', style='List Bullet')
document.add_paragraph('• Guardado local de ubicaciones favoritas o frecuentes de manera anónima mediante LocalStorage en el navegador.', style='List Bullet')

p = document.add_paragraph()
p.add_run('\nVersionado sugerido: ').bold = True
p.add_run('v0.1 (UI base, mapa renderizado y consumo de rutas estáticas), v0.2 (Controles interactivos de Alpha y parámetros de enrutamiento), v0.3 (Soporte dinámico multi-ciudad y panel analítico robusto), v1.0 (MVP del FrontEnd completo para validación experimental).')

doc_path = os.path.join(os.getcwd(), 'Implementacion_Frontend.docx')
document.save(doc_path)
print(f"Documento DOCX generado exitosamente en: {doc_path}")
